"""Score a filled blind labelling sheet against the mismatch agent and the auto-grade. $0, no model calls.

The sheet a reviewer fills in shows only protocol and registry evidence. The answer key (kept outside
the repo, never shown to the reviewer) holds the auto-grade category and the agent verdict per row.

    python -m scripts.score_labels --sheet FILLED.xlsx --key answer_key_v3.csv [--out report.md]

The sheet's "Label" tab needs columns "Row", "Trial", "What is compared", "Your label" and
"Note (optional)". Key columns: row_id, nct_id, field, auto_category, agent_verdict (plus optional
auto_reason, in_v2_set). Blank labels are skipped, so a part-filled sheet scores fine.
"""
from __future__ import annotations

import argparse
import collections
import csv
import json
from pathlib import Path

LABELS = ("Real difference", "Registry lagging", "Our reading wrong", "Can't tell")
AGENT_TO_LABEL = {"real_discrepancy": "Real difference", "registry_lagging": "Registry lagging",
                  "extraction_error": "Our reading wrong", "ambiguous": "Can't tell"}
# Same mapping as app.mismatch.agent.agreement: only settled auto-grade categories predict a label.
AUTO_TO_LABEL = {"caught_up": "Registry lagging", "partly_caught_up": "Registry lagging",
                 "never_changed": "Real difference", "changed_other_way": "Real difference",
                 "registry_exceeded_protocol": "Real difference"}
FIELD_FROM_SHEET = {"Number of sites": "sites", "Planned enrollment": "enrollment", "Phase": "phase",
                    "Randomized or not": "randomized", "Number of countries": "countries",
                    "Number of arms (groups)": "arms"}


def read_sheet(path: Path) -> list[dict]:
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True, data_only=True)["Label"]
    it = ws.iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else "" for h in next(it)]
    idx = {h: i for i, h in enumerate(head)}
    for need in ("Row", "Trial", "What is compared", "Your label"):
        if need not in idx:
            raise SystemExit(f"sheet is missing column {need!r}")
    out = []
    for r in it:
        if not r or r[idx["Row"]] in (None, ""):
            continue
        label = r[idx["Your label"]]
        label = str(label).strip() if label not in (None, "") else ""
        if label and label not in LABELS:
            raise SystemExit(f"row {r[idx['Row']]}: unknown label {label!r}; use one of {LABELS}")
        note = r[idx["Note (optional)"]] if "Note (optional)" in idx else None
        out.append({"row_id": str(r[idx["Row"]]), "nct_id": str(r[idx["Trial"]]),
                    "field": FIELD_FROM_SHEET.get(str(r[idx["What is compared"]]), str(r[idx["What is compared"]])),
                    "label": label, "note": "" if note is None else str(note)})
    return out


def read_key(path: Path) -> dict[str, dict]:
    with open(path, newline="") as f:
        return {r["row_id"]: r for r in csv.DictReader(f)}


def kappa(pairs: list[tuple[str, str]]) -> float | None:
    n = len(pairs)
    if n == 0:
        return None
    po = sum(a == b for a, b in pairs) / n
    ca, cb = collections.Counter(a for a, _ in pairs), collections.Counter(b for _, b in pairs)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / (n * n)
    return None if pe == 1 else round((po - pe) / (1 - pe), 3)


def rate(hits: int, n: int) -> dict:
    return {"agree": hits, "n": n, "rate": round(hits / n, 3) if n else None}


def score(sheet_rows: list[dict], key: dict[str, dict]) -> dict:
    rows = []
    for s in sheet_rows:
        k = key.get(s["row_id"])
        if k is None:
            raise SystemExit(f"row {s['row_id']} is not in the answer key")
        if (k["nct_id"], k["field"]) != (s["nct_id"], s["field"]):
            raise SystemExit(f"row {s['row_id']}: sheet says {s['nct_id']} {s['field']}, key says {k['nct_id']} {k['field']}")
        if not s["label"]:
            continue
        rows.append({**s, "agent": AGENT_TO_LABEL[k["agent_verdict"]], "auto_category": k["auto_category"],
                     "auto": AUTO_TO_LABEL.get(k["auto_category"]), "auto_reason": k.get("auto_reason", ""),
                     "in_v2_set": str(k.get("in_v2_set", "")).lower() == "true"})

    def block(rs):
        a = [(r["label"], r["agent"]) for r in rs]
        decided = [r for r in rs if r["label"] != "Can't tell"]
        settled = [r for r in rs if r["auto"]]
        settled_decided = [r for r in settled if r["label"] != "Can't tell"]
        return {"labelled": len(rs),
                "agent": {**rate(sum(x == y for x, y in a), len(a)), "kappa": kappa(a),
                          "excluding_cant_tell": rate(sum(r["label"] == r["agent"] for r in decided), len(decided))},
                "auto_grade_settled": {**rate(sum(r["label"] == r["auto"] for r in settled), len(settled)),
                                       "kappa": kappa([(r["label"], r["auto"]) for r in settled]),
                                       "excluding_cant_tell": rate(sum(r["label"] == r["auto"] for r in settled_decided),
                                                                   len(settled_decided))}}

    doubtful = [r for r in rows if r["auto_category"] == "ambiguous" and "doubtful" in r["auto_reason"]]
    finished = [r for r in rows if r["auto_category"] == "ambiguous" and "doubtful" not in r["auto_reason"]]
    return {
        "rows_in_sheet": len(sheet_rows), "labelled": len(rows),
        "trials_labelled": len({r["nct_id"] for r in rows}),
        "labels": dict(collections.Counter(r["label"] for r in rows)),
        "overall": block(rows),
        "v2_comparable_rows": block([r for r in rows if r["in_v2_set"]]),
        "by_field": {f: block([r for r in rows if r["field"] == f]) for f in sorted({r["field"] for r in rows})},
        "auto_ambiguous": {
            "doubtful_reading": {"n": len(doubtful), "labels": dict(collections.Counter(r["label"] for r in doubtful))},
            "finished_before_next_snapshot": {"n": len(finished),
                                              "labels": dict(collections.Counter(r["label"] for r in finished))}},
        "confusion_vs_agent": confusion(rows, "agent", LABELS),
        "confusion_vs_auto_category": confusion(rows, "auto_category", sorted({r["auto_category"] for r in rows})),
        "agent_disagrees_with_reviewer": [f"{r['row_id']} {r['nct_id']} {r['field']}: reviewer {r['label']}, "
                                          f"agent {r['agent']}" for r in rows if r["label"] != r["agent"]],
    }


def confusion(rows, col, cols) -> dict:
    t = {lab: {c: 0 for c in cols} for lab in LABELS}
    for r in rows:
        t[r["label"]][r[col]] += 1
    return t


def pct(x: dict) -> str:
    return "n/a (0 rows)" if not x["n"] else f"{x['agree']}/{x['n']} ({x['rate']:.0%})"


def report(res: dict) -> str:
    o = res["overall"]
    L = ["# Reviewer labels vs agent and auto-grade", "",
         f"Labelled {res['labelled']} of {res['rows_in_sheet']} rows ({res['trials_labelled']} trials). "
         f"Labels: " + ", ".join(f"{k} {v}" for k, v in sorted(res["labels"].items())), "",
         "| | Agent | Auto-grade (settled rows only) |", "|---|---|---|",
         f"| Agreement | {pct(o['agent'])} | {pct(o['auto_grade_settled'])} |",
         f"| Agreement, reviewer 'Can't tell' left out | {pct(o['agent']['excluding_cant_tell'])} | "
         f"{pct(o['auto_grade_settled']['excluding_cant_tell'])} |",
         f"| Cohen's kappa | {o['agent']['kappa']} | {o['auto_grade_settled']['kappa']} |",
         f"| Rows from the v2 set only | {pct(res['v2_comparable_rows']['agent'])} | "
         f"{pct(res['v2_comparable_rows']['auto_grade_settled'])} |", "",
         "## Per field", "", "| Field | Labelled | Agent | Auto-grade (settled) |", "|---|---|---|---|"]
    for f, b in res["by_field"].items():
        L.append(f"| {f} | {b['labelled']} | {pct(b['agent'])} | {pct(b['auto_grade_settled'])} |")
    aa = res["auto_ambiguous"]
    L += ["", "## Rows the auto-grade could not settle", "",
          f"- Our reading flagged doubtful ({aa['doubtful_reading']['n']}): " +
          ", ".join(f"{k} {v}" for k, v in sorted(aa["doubtful_reading"]["labels"].items())),
          f"- Trial finished before the next snapshot ({aa['finished_before_next_snapshot']['n']}): " +
          ", ".join(f"{k} {v}" for k, v in sorted(aa["finished_before_next_snapshot"]["labels"].items())), ""]
    for name, col in (("agent", "confusion_vs_agent"), ("auto-grade category", "confusion_vs_auto_category")):
        t = res[col]
        cols = list(next(iter(t.values())).keys())
        L += [f"## Confusion: reviewer (rows) vs {name} (columns)", "",
              "| Reviewer | " + " | ".join(cols) + " |", "|---|" + "---|" * len(cols)]
        L += [f"| {lab} | " + " | ".join(str(t[lab][c]) for c in cols) + " |" for lab in LABELS]
        L.append("")
    L += ["## Where the reviewer and the agent differ", ""] + [f"- {x}" for x in res["agent_disagrees_with_reviewer"]]
    return "\n".join(L) + "\n"


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--sheet", type=Path, required=True)
    ap.add_argument("--key", type=Path, required=True)
    ap.add_argument("--out", type=Path, help="write the markdown report here (and .json beside it)")
    a = ap.parse_args()
    res = score(read_sheet(a.sheet), read_key(a.key))
    md = report(res)
    if a.out:
        a.out.write_text(md)
        a.out.with_suffix(".json").write_text(json.dumps(res, indent=1))
    print(md)


if __name__ == "__main__":
    main()
